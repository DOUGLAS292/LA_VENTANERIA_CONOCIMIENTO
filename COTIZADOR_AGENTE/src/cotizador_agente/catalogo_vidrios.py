"""Catálogo maestro de vidrios de La Ventanería, leído en tiempo de ejecución.

El archivo vive en Drive (lleva precios y proveedores, es CONFIDENCIAL) y nunca en
GitHub. Este módulo lee SOLO las columnas técnicas: ningún precio entra al agente,
porque los precios los pone el software de cotización.

El agente lo usa así: calcula con la NSR-10 el espesor mínimo por tipo de vidrio
(`vidrio_por_viento`) y si se exige vidrio de seguridad (`vidrio_seguridad`), y
luego pide al catálogo los vidrios que cumplen y caben en la serie escogida.
"""

import json
import os
from dataclasses import asdict, dataclass
from pathlib import Path

from .criterio_tecnico import DATOS

ENV_RUTA = "LV_CATALOGO_VIDRIOS"
HOJA = "Catálogo"

# Columnas que se leen (por nombre de encabezado). Las de precio no se leen nunca.
COLUMNAS = {
    "codigo": "Código", "descripcion": "Descripción", "familia": "Familia", "tipo_nsr": "Tipo NSR-10",
    "l1": "Lámina 1 (mm)", "l2": "Lámina 2 (mm)", "l3": "Lámina 3 (mm)", "l4": "Lámina 4 (mm)",
    "interlayer": "Interlayer (mm)", "camara": "Cámara (mm)", "color": "Color / tono", "prestacion": "Prestación",
    "estado": "Estado", "por_revisar": "Por revisar",
}
FAMILIAS_SEGURIDAD = {"Templado", "Laminado", "Laminado templado", "Blindado"}
FAMILIAS_INGENIERIA = {"Insulado (DVH)", "Blindado"}
FAMILIAS_NO_VENTANA = {"Espejo", "Otro relleno", "Sin vidrio"}


@dataclass
class VidrioCatalogo:
    codigo: str
    descripcion: str
    familia: str
    tipo_nsr: str | None
    espesor_nominal_mm: float | None
    espesor_total_mm: float | None
    color: str | None
    prestacion: str | None
    estado: str
    por_revisar: str | None

    @property
    def seguridad(self) -> bool:
        return self.familia in FAMILIAS_SEGURIDAD


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def cargar(ruta: Path | str) -> list[VidrioCatalogo]:
    """Lee la hoja 'Catálogo'. Calcula los espesores desde las láminas (no depende de fórmulas de Excel)."""
    from openpyxl import load_workbook

    ws = load_workbook(ruta, read_only=True, data_only=True)[HOJA]
    filas = ws.iter_rows(values_only=True)
    encabezado = [str(c).strip() if c is not None else "" for c in next(filas)]
    idx = {}
    for clave, nombre in COLUMNAS.items():
        if nombre in encabezado:
            idx[clave] = encabezado.index(nombre)
    faltan = [COLUMNAS[k] for k in ("codigo", "descripcion", "familia") if k not in idx]
    if faltan:
        raise ValueError(f"Al catálogo le faltan columnas: {', '.join(faltan)}")

    vidrios = []
    for fila in filas:
        val = lambda k: fila[idx[k]] if k in idx and idx[k] < len(fila) else None  # noqa: E731
        codigo, familia = val("codigo"), val("familia")
        if codigo in (None, "") or familia in (None, ""):
            continue
        laminas = [x for x in (_num(val(k)) for k in ("l1", "l2", "l3", "l4")) if x is not None]
        nominal = sum(laminas) if laminas else None
        total = None if nominal is None else round(nominal + (_num(val("interlayer")) or 0) + (_num(val("camara")) or 0), 2)
        vidrios.append(VidrioCatalogo(
            codigo=str(codigo).strip(), descripcion=str(val("descripcion") or ""), familia=str(familia),
            tipo_nsr=val("tipo_nsr"), espesor_nominal_mm=nominal, espesor_total_mm=total,
            color=val("color"), prestacion=val("prestacion"), estado=str(val("estado") or "Activo"),
            por_revisar=val("por_revisar"),
        ))
    return vidrios


def rango_serie(codigo_sistema: str) -> tuple[float, float] | None:
    """Espesor de vidrio admisible de una serie, según su ficha (datos/sistemas)."""
    ruta = DATOS / "sistemas" / f"{codigo_sistema}.json"
    if not ruta.exists():
        return None
    c = json.loads(ruta.read_text(encoding="utf-8")).get("caracteristicas", {})
    for clave in ("vidrio_mm", "vidrio_corrediza_mm", "vidrios_admisibles_mm"):
        if clave in c and c[clave]:
            return float(min(c[clave])), float(max(c[clave]))
    return None


def opciones(vidrios: list[VidrioCatalogo], sistema: str, minimos_mm: dict[str, float],
             exige_seguridad: bool = False, color: str | None = None, prestacion: str | None = None) -> dict:
    """Vidrios del catálogo que cumplen la NSR-10 y caben en la serie.

    minimos_mm: espesor mínimo por tipo NSR, tal como lo da `vidrio_por_viento`, p. ej.
    {"recocido": 6, "templado": 4, "laminado_pvb": 5}. Un tipo ausente o None no se ofrece.
    """
    rango = rango_serie(sistema)
    cumplen, a_ingenieria = [], []
    for v in vidrios:
        if v.estado == "Inactivo" or v.familia in FAMILIAS_NO_VENTANA:
            continue
        if v.familia in FAMILIAS_INGENIERIA:
            a_ingenieria.append(v.codigo)
            continue
        minimo = minimos_mm.get(v.tipo_nsr or "")
        if minimo is None or v.espesor_nominal_mm is None or v.espesor_nominal_mm < minimo:
            continue
        if exige_seguridad and not v.seguridad:
            continue
        if rango and not rango[0] <= v.espesor_nominal_mm <= rango[1]:
            continue
        if color and color.lower() not in (v.color or "").lower():
            continue
        if prestacion and prestacion.lower() not in (v.prestacion or "").lower():
            continue
        d = asdict(v)
        d["seguridad"] = v.seguridad
        cumplen.append(d)
    cumplen.sort(key=lambda d: (d["espesor_nominal_mm"], d["codigo"]))
    return {
        "sistema": sistema,
        "rango_serie_mm": rango,
        "opciones": cumplen,
        "insulados_y_blindados_a_ingenieria": a_ingenieria,
        "nota": "Sin precios: el software de cotización los pone. Las opciones con 'por_revisar' tienen datos por confirmar.",
    }


def ruta_por_defecto() -> Path | None:
    ruta = os.environ.get(ENV_RUTA)
    return Path(ruta) if ruta else None
